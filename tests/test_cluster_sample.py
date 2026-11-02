import csv
import json
from collections import Counter

import pytest

from scripts import cluster_sample as cs


def _clusters(sizes: dict[str, int]) -> list[str]:
    return [c for c, n in sizes.items() for _ in range(n)]


def _counts(selected: dict[int, str]) -> dict[str, int]:
    return dict(Counter(selected.values()))


def _config(**kwargs) -> cs.SampleConfig:
    return cs.SampleConfig(**{"tail_bins": (), "single_max_size": 0, **kwargs})


def test_remap_merges_tail_clusters_by_size_bins():
    clusters = _clusters({"a": 1, "b": 1, "c": 2, "d": 3, "e": 5, "f": 50})
    new = dict(zip(clusters, cs.remap_clusters(clusters, (1, 3))))
    assert new == {"a": "tail:1", "b": "tail:1", "c": "tail:2-3", "d": "tail:2-3", "e": "e", "f": "f"}


def test_remap_without_bins_keeps_ids():
    clusters = _clusters({"a": 1, "b": 4})
    assert cs.remap_clusters(clusters, ()) == clusters


def test_allocation_is_sublinear_and_hits_total():
    clusters = _clusters({"a": 1000, "b": 100, "c": 10})
    selected, _ = cs.sample(clusters, _config(total=100, rounding="largest_remainder"))
    assert _counts(selected) == {"a": 76, "b": 19, "c": 5}


def test_cluster_cap_redistributes():
    clusters = _clusters({"big": 1000, "mid": 40, "small": 30})
    counts = _counts(cs.sample(clusters, _config(total=60, alpha=1.0, cluster_cap=20, rounding="largest_remainder"))[0])
    assert counts["big"] == 20 and sum(counts.values()) == 60


def test_small_unmerged_clusters_get_exactly_one():
    # tail_bins=1,3 合并 1~3；4、5 未合并且 <= 5，固定抽 1 条；6 及以上正常分配
    sizes = {**{f"s{i}": 1 for i in range(30)}, **{f"m{i}": 3 for i in range(10)},
             "four": 4, "five": 5, "six": 6, "big": 200}
    config = cs.SampleConfig(total=150, tail_bins=(1, 3), single_max_size=5, rounding="largest_remainder")
    counts = _counts(cs.sample(_clusters(sizes), config)[0])
    assert counts["four"] == 1 and counts["five"] == 1
    assert counts["six"] > 1
    assert counts["tail:1"] > 1 and counts["tail:2-3"] > 1  # 合并后的长尾簇不受限
    assert sum(counts.values()) == 150


def test_small_clusters_guaranteed_even_when_budget_is_tight():
    # 300 个大小为 10 的簇 + 一个大簇；按权重分时小簇配额远小于 1，保底后每个都抽到 1 条
    sizes = {**{f"c{i}": 10 for i in range(300)}, "big": 100000}
    config = cs.SampleConfig(total=400, tail_bins=(), single_max_size=20)
    counts = _counts(cs.sample(_clusters(sizes), config)[0])
    assert all(counts[f"c{i}"] == 1 for i in range(300))
    assert counts["big"] == 100 and sum(counts.values()) == 400


def test_too_many_small_clusters_rejected():
    config = cs.SampleConfig(total=3, tail_bins=(), single_max_size=5)
    with pytest.raises(ValueError, match="保底"):
        cs.sample(_clusters({f"c{i}": 2 for i in range(5)}), config)


def test_single_max_size_zero_disables_rule():
    config = cs.SampleConfig(total=10, tail_bins=(), single_max_size=0, rounding="largest_remainder")
    assert _counts(cs.sample(_clusters({"a": 5, "b": 5}), config)[0]) == {"a": 5, "b": 5}
    config = cs.SampleConfig(total=10, tail_bins=(), single_max_size=5)
    assert _counts(cs.sample(_clusters({"a": 5, "b": 5}), config)[0]) == {"a": 1, "b": 1}


def test_capacity_shortfall_takes_everything():
    selected, _ = cs.sample(_clusters({"a": 3, "b": 2}), _config(total=100))
    assert len(selected) == 5


def test_pps_is_reproducible_and_exact():
    clusters = _clusters({f"c{i}": 3 + i % 17 for i in range(500)})
    config = _config(total=200, seed=7)
    first, second = cs.sample(clusters, config)[0], cs.sample(clusters, config)[0]
    assert first == second and len(first) == 200
    assert max(_counts(first).values()) <= 1


def test_pps_expected_counts_match_quota():
    sizes = {"x": 4, "y": 9, **{f"z{i}": 2 for i in range(20)}}
    config = _config(total=6)
    quota = cs.allocate(sizes, cs.capacities(sizes, config), config.total, config.alpha)
    hits = Counter()
    for seed in range(400):
        hits.update(_counts(cs.sample(_clusters(sizes), _config(total=6, seed=seed))[0]))
    for c, q in quota.items():
        assert hits[c] / 400 == pytest.approx(q, abs=0.08)


def test_selected_rows_match_detail():
    clusters = _clusters({f"c{i}": i + 1 for i in range(60)})
    selected, detail = cs.sample(clusters, cs.SampleConfig(total=150, tail_bins=(1, 3, 5), cluster_cap=8))
    assert len(selected) == 150
    assert {c: d["sampled"] for c, d in detail.items() if d["sampled"]} == _counts(selected)


@pytest.mark.parametrize("kwargs", [
    {"total": 0}, {"alpha": -1}, {"tail_bins": (3, 1)}, {"tail_bins": (1, 1)}, {"rounding": "x"},
    {"cluster_cap": 0}, {"single_max_size": -1}, {"id_column": "a", "cluster_column": "a"}, {"cluster_column": "new_cluster_id"}, {"cluster_column": "segment"},
])
def test_invalid_config_rejected(kwargs):
    with pytest.raises(ValueError):
        cs.SampleConfig(**kwargs)


def test_missing_column_and_reserved_output_column_rejected():
    config = cs.SampleConfig(id_column="uid", cluster_column="cid")
    with pytest.raises(ValueError, match="缺少字段"):
        cs.cluster_keys([{"uid": "1"}], config)
    with pytest.raises(ValueError, match="new_cluster_id"):
        cs.cluster_keys([{"uid": "1", "cid": "a", "new_cluster_id": "x"}], config)


def _write_csv(path, header, rows):
    with path.open("w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        writer.writerows(rows)


def _read_csv(path):
    with path.open(newline="") as f:
        reader = csv.DictReader(f)
        return reader.fieldnames, list(reader)


def test_cli_defaults_and_report(tmp_path):
    src, out, report = tmp_path / "batch.csv", tmp_path / "out.csv", tmp_path / "report.json"
    _write_csv(src, ["id", "cluster_id"], [(k, c) for k, c in enumerate(_clusters({"a": 30, "b": 5, "c": 1, "d": 1}))])
    assert cs.main(["--input", str(src), "--output", str(out), "--total", "10", "--report", str(report)]) == 0
    header, rows = _read_csv(out)
    assert header == ["id", "cluster_id", "new_cluster_id", "segment", "cluster_size"] and len(rows) == 10
    assert {(r["cluster_id"], r["segment"], r["cluster_size"]) for r in rows} >= {("a", "normal:>=6", "30"), ("b", "small:4-5", "5")}
    data = json.loads(report.read_text())
    assert data["config"]["tail_bins"] == [1, 3]
    assert set(data["clusters"]) == {"a", "b", "tail:1"} and data["clusters"]["tail:1"]["size"] == 2
    segs = data["segments"]
    assert {k: (v["clusters"], v["size"]) for k, v in segs.items()} == {
        "normal:>=6": (1, 30), "small:4-5": (1, 5), "tail:1": (1, 2)}
    assert segs["small:4-5"]["sampled"] == 1 and sum(v["sampled"] for v in segs.values()) == 10


def test_cli_empty_tail_bins_disables_merge(tmp_path):
    src, out, report = tmp_path / "batch.csv", tmp_path / "out.csv", tmp_path / "report.json"
    _write_csv(src, ["id", "cluster_id"], [(k, c) for k, c in enumerate(_clusters({"a": 30, "c": 1}))])
    assert cs.main(["--input", str(src), "--output", str(out), "--total", "5", "--tail-bins", "",
                    "--report", str(report)]) == 0
    assert set(json.loads(report.read_text())["clusters"]) == {"a", "c"}


@pytest.mark.parametrize("text, expected", [("1,3,5", [1, 3, 5]), ("1, 3 ,5", [1, 3, 5]), ("2", [2]), ("", [])])
def test_tail_bins_comma_format(text, expected):
    assert cs._int_list(text) == expected


@pytest.mark.parametrize("text", ["1 3", "1,,3", "a"])
def test_tail_bins_bad_format_rejected(text, tmp_path):
    with pytest.raises(SystemExit):
        cs.parse_args(["--input", "x.csv", "--output", "y.csv", "--tail-bins", text])


def test_cli_custom_columns_keep_other_fields_and_allow_duplicate_ids(tmp_path):
    src, out = tmp_path / "batch.csv", tmp_path / "out.csv"
    topics = ["x"] * 20 + ["y"] * 4 + ["z"]
    # query 列故意重复：同一个 query 出现在多行，每行仍独立参与抽样
    original = [(f"q{k % 5}", t, f"text, {k}", str(k)) for k, t in enumerate(topics)]
    _write_csv(src, ["query", "topic", "raw_text", "row"], original)
    assert cs.main(["--input", str(src), "--output", str(out), "--id-column", "query",
                    "--cluster-column", "topic", "--total", "25", "--tail-bins", "1", "--single-max-size", "0"]) == 0
    header, rows = _read_csv(out)
    assert header == ["query", "topic", "raw_text", "row", "new_cluster_id", "segment", "cluster_size"]
    assert [tuple(r[k] for k in header[:4]) for r in rows] == original  # 全部抽中时与输入逐行一致
    assert [r["new_cluster_id"] for r in rows] == [("tail:1" if t == "z" else t) for t in topics]


def test_cli_jsonl_keeps_nested_fields(tmp_path):
    src, out = tmp_path / "batch.jsonl", tmp_path / "out.jsonl"
    records = [{"uid": k, "cid": k % 3, "meta": {"text": f"t{k}", "tags": [k]}} for k in range(30)]
    src.write_text("".join(json.dumps(r) + "\n" for r in records))
    assert cs.main(["--input", str(src), "--output", str(out), "--id-column", "uid", "--cluster-column", "cid",
                    "--total", "6"]) == 0
    rows = [json.loads(line) for line in out.read_text().splitlines()]
    assert len(rows) == 6
    for r in rows:
        assert {k: v for k, v in r.items() if k not in cs.OUTPUT_COLUMNS} == records[r["uid"]]
        assert r["new_cluster_id"] == str(r["cid"])


def test_cli_xlsx_roundtrip(tmp_path):
    from openpyxl import Workbook, load_workbook
    src, out = tmp_path / "batch.xlsx", tmp_path / "out.xlsx"
    book = Workbook()
    book.active.append(["id", "cluster_id", "turns"])
    for k, c in enumerate(_clusters({"a": 20, "b": 1})):
        book.active.append([f"q{k}", c, '["x"]'])
    book.save(src)
    assert cs.main(["--input", str(src), "--output", str(out), "--total", "5", "--tail-bins", ""]) == 0
    rows = list(load_workbook(out).active.values)
    assert rows[0] == ("id", "cluster_id", "turns", "new_cluster_id", "segment", "cluster_size")
    assert len(rows) == 6 and all(r[2] == '["x"]' for r in rows[1:])


@pytest.mark.parametrize("bins, single, expected", [
    ((1, 3), 5, ("small:4-5", "normal:>=6")),
    ((1, 3, 5), 8, ("small:6-8", "normal:>=9")),
    ((1, 3), 4, ("small:4", "normal:>=5")),
    ((1, 3), 3, (None, "normal:>=4")),
    ((), 0, (None, "normal:>=1")),
    ((), 2, ("small:1-2", "normal:>=3")),
])
def test_segment_names_follow_config(bins, single, expected):
    assert cs.segment_names(cs.SampleConfig(tail_bins=bins, single_max_size=single)) == expected


def test_segments_follow_custom_bins():
    clusters = _clusters({"a": 1, "b": 2, "c": 4, "d": 5, "e": 7, "f": 9, "g": 30})
    _, detail = cs.sample(clusters, cs.SampleConfig(total=20, tail_bins=(1, 3, 5), single_max_size=8))
    assert {c: d["segment"] for c, d in detail.items()} == {
        "tail:1": "tail:1", "tail:2-3": "tail:2-3", "tail:4-5": "tail:4-5",
        "e": "small:6-8", "f": "normal:>=9", "g": "normal:>=9"}


def test_fine_column_splits_only_big_clusters():
    # big 30 条拆成 big.0(20)/big.1(9)/big.2(1)，每个细簇 1 条；mid 5 条照旧 1 条；剩余名额给 tail
    clusters = _clusters({"big": 30, "mid": 5, "t1": 1, "t2": 1, "t3": 1})
    fine = ["big.0"] * 20 + ["big.1"] * 9 + ["big.2"] + ["mid.0"] * 3 + ["mid.1"] * 2 + ["t1.0", "t2.0", "t3.0"]
    config = cs.SampleConfig(total=6, tail_bins=(1,), single_max_size=20)
    selected, detail = cs.sample(clusters, config, fine)
    assert _counts(selected) == {"big.0": 1, "big.1": 1, "big.2": 1, "mid": 1, "tail:1": 2}
    assert {c: d["segment"] for c, d in detail.items()} == {
        "big.0": "split:>=21", "big.1": "split:>=21", "big.2": "split:>=21", "mid": "small:2-20", "tail:1": "tail:1"}


def test_fine_column_must_nest():
    with pytest.raises(ValueError, match="嵌套"):
        cs.sample(["a", "a", "b"], cs.SampleConfig(total=2, tail_bins=(), single_max_size=0), ["x", "y", "x"])


def test_cli_fine_column(tmp_path):
    src, out = tmp_path / "batch.csv", tmp_path / "out.csv"
    rows = [[f"q{k}", c, f"{c}.{k % 3}"] for k, c in enumerate(_clusters({"a": 9, "b": 1}))]
    _write_csv(src, ["id", "cluster_id", "fine_cluster_id"], rows)
    assert cs.main(["--input", str(src), "--output", str(out), "--total", "4", "--tail-bins", "1",
                    "--single-max-size", "5", "--fine-column", "fine_cluster_id"]) == 0
    _, got = _read_csv(out)
    assert sorted(r["new_cluster_id"] for r in got) == ["a.0", "a.1", "a.2", "tail:1"]
    assert {r["cluster_size"] for r in got if r["new_cluster_id"] != "tail:1"} == {"9"}
